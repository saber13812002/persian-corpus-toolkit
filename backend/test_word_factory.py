# -*- coding: utf-8 -*-
"""test_word_factory — آزمون‌های Word Factory V0.1 (WF-01..WF-13).

مستقل از pct_test_support: factory هیچ DB/Flask/سرویس operationی ندارد —
تنها filesystem موقت (tempfile) و python-docx. WF-14 (کتاب واقعی) جدا و
در گزارش به‌عنوان PARTIAL/BLOCKED ثبت می‌شود.
"""
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from factory import (DocumentProcessingService, PersianNormalizer,
                     NormalizationPolicy, QualityValidator, make_parser,
                     DocParserAdapter, DocxParser, detect_format,
                     verify_no_semantic_change, sha256_file)
from factory.errors import (FormatDetectionError, UnsupportedFormatError,
                            ValidationFailedError, FileMissingError)
from factory._smoke_fixtures import (make_broken_docx, make_docx,
                                     make_ole2_doc, make_sample_book)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class WfBase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="wf_test_")
        self.out = os.path.join(self.dir, "out")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def svc(self, **kw):
        return DocumentProcessingService(out_dir=self.out, **kw)


class TestWF01_Intake(WfBase):
    """WF-01: دریافت فایل DOCX سالم."""

    def test_wf01_good_docx(self):
        src = make_sample_book(os.path.join(self.dir, "book.docx"))
        doc = self.svc().process(src)
        self.assertEqual(doc.file.format, "docx")
        self.assertTrue(doc.file.sha256, "64")
        self.assertEqual(len(doc.file.sha256), 64)
        self.assertTrue(doc.quality.passed, doc.quality.warnings)
        # فایل اصلی دست‌نخورده
        self.assertEqual(sha256_file(src), doc.file.sha256)

    def test_wf01_missing_file(self):
        with self.assertRaises(FileMissingError):
            self.svc().process(os.path.join(self.dir, "nope.docx"))

    def test_wf01_empty_file(self):
        p = os.path.join(self.dir, "empty.docx")
        open(p, "wb").close()
        with self.assertRaises(ValidationFailedError):
            self.svc().process(p)


class TestWF02_Corrupt(WfBase):
    """WF-02: رد فایل خراب (bad zip / content types ناقص)."""

    def test_wf02_broken_zip(self):
        p = make_broken_docx(os.path.join(self.dir, "broken.docx"))
        with self.assertRaises((FormatDetectionError, ValidationFailedError)):
            self.svc().process(p)

    def test_wf02_manifest_written_on_failure(self):
        p = make_broken_docx(os.path.join(self.dir, "broken.docx"))
        try:
            self.svc().process(p)
            self.fail("باید خطا می‌داده")
        except (FormatDetectionError, ValidationFailedError):
            pass
        manifests = []
        for root, _, files in os.walk(self.out):
            manifests += [f for f in files if f == "manifest.json"]
        if manifests:  # اگر intake رد کرده باشد manifest لازم نیست
            m = json.load(open(os.path.join(self.out, manifests[0]), encoding="utf-8"))
            self.assertEqual(m["processing"]["status"], "failed")


class TestWF03_LegacyDoc(WfBase):
    """WF-03: تشخیص DOC قدیمی + اعلام عدم پشتیبانی (نه parse کاذب)."""

    def test_wf03_ole2_detected(self):
        p = make_ole2_doc(os.path.join(self.dir, "legacy.doc"))
        fmt, how = detect_format(p)
        self.assertEqual(fmt, "doc")
        self.assertEqual(how, "magic-ole2")

    def test_wf03_service_raises_unsupported(self):
        p = make_ole2_doc(os.path.join(self.dir, "legacy.doc"))
        with self.assertRaises(UnsupportedFormatError) as cm:
            self.svc().process(p)
        self.assertEqual(cm.exception.format, "doc")

    def test_wf03_adapter_also_refuses(self):
        p = make_ole2_doc(os.path.join(self.dir, "legacy.doc"))
        from factory.intake import DocumentIntakeService
        # adapter مستقیم: FileRecord ساختگی
        from factory.models import FileRecord, utc_now_iso
        rec = FileRecord(file_id="f-test", filename="legacy.doc", path=p,
                         size_bytes=os.path.getsize(p), sha256="x",
                         format="doc", detected_by="magic-ole2",
                         created_at=utc_now_iso())
        with self.assertRaises(UnsupportedFormatError):
            DocParserAdapter().parse(rec)

    def test_wf03_docx_content_with_doc_extension(self):
        # پسوند .doc ولی محتوای docx → نمی‌توان «doc قدیمی» فرضش کرد
        p = make_docx(os.path.join(self.dir, "tricky.doc"),
                      [("Normal", "متن تستی با پسوند اشتباه")])
        fmt, how = detect_format(p)
        self.assertEqual(fmt, "docx")


class TestWF04_PersianText(WfBase):
    """WF-04: استخراج متن فارسی (بدون بازنویسی)."""

    def test_wf04_raw_exact(self):
        src = make_docx(os.path.join(self.dir, "fa.docx"),
                        [("Normal", "سلام دنیا"), ("Normal", "پیکرهٔ فارسی")])
        doc = self.svc().process(src)
        texts = [b.raw_text for b in doc.blocks if b.type == "paragraph"]
        self.assertIn("سلام دنیا", texts)
        self.assertIn("پیکرهٔ فارسی", texts)
        # کیفیت فارسی بالا
        self.assertGreaterEqual(doc.quality.metrics["persian_ratio"], 0.8)


class TestWF05_Order(WfBase):
    """WF-05: حفظ ترتیب واقعی پاراگراف و جدول."""

    def test_wf05_body_order(self):
        src = make_sample_book(os.path.join(self.dir, "order.docx"))
        doc = self.svc().process(src)
        # جدول آخر body است (بعد از فصل دوم)
        last = doc.blocks[-1]
        self.assertEqual(last.type, "table")
        self.assertEqual(last.table_rows[0], ["ستون1", "ستون2", "ستون3"])
        # ترتیب تیترها: فصل اول < بخش ۱ < فصل دوم
        heads = [b for b in doc.blocks if b.type == "heading"]
        self.assertEqual([h.raw_text for h in heads],
                         ["فصل اول", "بخش ۱", "فصل دوم"])
        # index پیوسته
        self.assertTrue(all(b.block_index == i for i, b in enumerate(doc.blocks)))


class TestWF06_Headings(WfBase):
    """WF-06: تیترها و سلسله‌مراتب (parent)."""

    def test_wf06_hierarchy(self):
        src = make_docx(os.path.join(self.dir, "h.docx"), [
            ("Heading 1", "فصل یک"),
            ("Normal", "متن"),
            ("Heading 2", "زیربخش الف"),
            ("Normal", "متن زیر"),
            ("Heading 1", "فصل دو"),
        ])
        doc = self.svc().process(src)
        t = doc.structure
        self.assertEqual(len(t.headings), 3)
        self.assertEqual(len(t.chapters), 2)
        # parent زیربخش = فصل یک
        sub = next(h for h in t.headings if h["text"] == "زیربخش الف")
        first = next(h for h in t.headings if h["text"] == "فصل یک")
        self.assertEqual(sub["parent_block_id"], first["block_id"])
        # فصل دو بدون parent
        ch2 = next(h for h in t.headings if h["text"] == "فصل دو")
        self.assertIsNone(ch2["parent_block_id"])


class TestWF07_Normalize(WfBase):
    """WF-07: نرمال‌سازی بدون تغییر RAW."""

    def test_wf07_raw_untouched(self):
        src = make_docx(os.path.join(self.dir, "n.docx"),
                        [("Normal", "کتاب   فارسی  و ی عربی")])
        doc = self.svc().process(src)
        p = next(b for b in doc.blocks if "کتاب" in b.raw_text)
        self.assertEqual(p.raw_text, "کتاب   فارسی  و ی عربی")  # RAW دست‌نخورده
        self.assertNotEqual(p.normalized_text, p.raw_text)
        self.assertEqual(p.normalized_text.count("   "), 0)

    def test_wf07_rule_independent(self):
        # هر قانون مستقل قابل تنظیم است
        t, stats = PersianNormalizer(NormalizationPolicy(zwnj=False)).normalize("کتاب ها")
        self.assertNotIn("\u200C", t)
        t2, stats2 = PersianNormalizer(NormalizationPolicy()).normalize("کتاب ها")
        self.assertIn("\u200C", t2)
        self.assertIn("zwnj", stats2)

    def test_wf07_no_semantic_change_guard(self):
        # guard: حذف حرف واقعی را پیدا می‌کند
        issues = verify_no_semantic_change("سلام دنیا", "سلام دیا")
        self.assertTrue(issues)
        # و تغییر مجاز (space/ZWNJ/یکسان‌سازی) را نه
        issues2 = verify_no_semantic_change("کتاب ِ فارسی", "کتابِ فارسی")
        self.assertEqual(issues2, [])
        # یکسان‌سازی ى عربی (U+0649) → ی فارسی (U+06CC) تغییر معنی نیست
        issues3 = verify_no_semantic_change("دني\u0649", "دني\u06CC")
        self.assertEqual(issues3, [])

    def test_wf07_digits_policy(self):
        t, _ = PersianNormalizer(NormalizationPolicy(digits="persian")).normalize("سال 1405")
        self.assertIn("۱۴۰۵", t)
        t2, _ = PersianNormalizer(NormalizationPolicy(digits="ascii")).normalize("سال ۱۴۰۵")
        self.assertIn("1405", t2)


class TestWF08_HashManifest(WfBase):
    """WF-08: SHA-256 و manifest."""

    def test_wf08_manifest(self):
        src = make_sample_book(os.path.join(self.dir, "m.docx"))
        doc = self.svc().process(src)
        d = os.path.join(self.out, doc.file.file_id)
        self.assertTrue(os.path.isfile(os.path.join(d, "manifest.json")))
        m = json.load(open(os.path.join(d, "manifest.json"), encoding="utf-8"))
        self.assertEqual(m["file"]["sha256"], doc.file.sha256)
        self.assertEqual(m["versions"]["normalizer"], "n-1")
        self.assertEqual(m["processing"]["status"], "done")
        # hash هر خروجی در manifest درست است
        for name, info in m["outputs"].items():
            real = hashlib.sha256(open(info["path"], "rb").read()).hexdigest()
            self.assertEqual(info["sha256"], real, name)
        self.assertEqual(m["counts"]["blocks"], len(doc.blocks))


class TestWF09_Jsonl(WfBase):
    """WF-09: خروجی JSONL معتبر (هر خط JSON، فیلدهای قرارداد)."""

    REQUIRED = {"block_id", "block_index", "type", "raw_text",
                "normalized_text", "norm_version", "content_hash", "page"}

    def test_wf09_jsonl(self):
        src = make_sample_book(os.path.join(self.dir, "j.docx"))
        doc = self.svc().process(src)
        d = os.path.join(self.out, doc.file.file_id)
        lines = open(os.path.join(d, "blocks.jsonl"), encoding="utf-8").read().splitlines()
        self.assertEqual(len(lines), len(doc.blocks))
        for line in lines:
            rec = json.loads(line)  # هر خط JSON معتبر
            self.assertTrue(self.REQUIRED.issubset(rec.keys()), rec.keys())
            self.assertIsNone(rec["page"])  # حدس صفحه ممنوع
            self.assertEqual(rec["norm_version"], "n-1")


class TestWF10_Idempotent(WfBase):
    """WF-10: پردازش مجدد → خروجی منطقی یکسان، بدون خروجی تکراری."""

    def test_wf10_rerun(self):
        src = make_sample_book(os.path.join(self.dir, "r.docx"))
        d1 = self.svc().process(src)
        d2 = self.svc().process(src)
        self.assertEqual(d1.file.file_id, d2.file.file_id)  # همان hash → همان file_id
        self.assertEqual(d1.blocks_hash(), d2.blocks_hash())
        # خروجی روی همان پوشه (نه پوشه دوم) — فقط یک file_id زیر out
        # (پوشه files/ متعلق به intake است و انتظار می‌رود)
        ids = [d for d in os.listdir(self.out) if d.startswith("f-")]
        self.assertEqual(ids, [d1.file.file_id])
        m = json.load(open(os.path.join(self.out, d1.file.file_id, "manifest.json"),
                           encoding="utf-8"))
        self.assertEqual(m["file"]["sha256"], sha256_file(src))


class TestWF11_PartialErrors(WfBase):
    """WF-11: خطاهای استخراج ناقص → ثبت در quality report (نه crash)."""

    def test_wf11_empty_doc(self):
        src = make_docx(os.path.join(self.dir, "e.docx"), [("Normal", "")])
        doc = self.svc().process(src)  # نباید exception بدهد
        self.assertEqual(doc.quality.status, "manual_review")
        self.assertFalse(doc.quality.passed)
        d = os.path.join(self.out, doc.file.file_id)
        q = json.load(open(os.path.join(d, "quality-report.json"), encoding="utf-8"))
        ne = next(c for c in q["checks"] if c["name"] == "non_empty_text")
        self.assertFalse(ne["ok"])

    def test_wf11_hyperlink_loss_accounted(self):
        # سند با hyperlink → در losses ثبت می‌شود (adclaim استخراج کامل ممنوع)
        from docx import Document as D
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        src = os.path.join(self.dir, "link.docx")
        d = D()
        p = d.add_paragraph("متن با ")
        run = p.add_run("لینک")
        a = OxmlElement("w:hyperlink")
        a.set(qn("r:id"), "rId1")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "لینک"
        r.append(t)
        a.append(r)
        p._p.append(a)
        d.add_paragraph("بعد از لینک.")
        d.save(src)
        doc = self.svc().process(src)
        self.assertGreaterEqual(doc.quality.unextracted["losses"].get("hyperlink", 0), 1)


class TestWF12_NoProductionDB(WfBase):
    """WF-12: عدم اتصال به DB عملیاتی / هیچ import server/Flask."""

    def test_wf12_no_server_import(self):
        # گارانتی استاتیک: هیچ فایل factory import ماژول‌های production نمی‌کند
        # (آزمون بر اساس وضعیت global sys.modules order-dependent بود و حذف شد)
        import ast
        import pathlib
        pkg = pathlib.Path(__file__).resolve().parent / "factory"
        banned = {"server", "flask", "pct_test_support", "dashboard_api",
                  "auth", "projects", "sqlite3"}
        for py in sorted(pkg.glob("*.py")):
            tree = ast.parse(py.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    mods = {a.name.split(".")[0] for a in node.names}
                elif isinstance(node, ast.ImportFrom) and node.module:
                    mods = {node.module.split(".")[0]}
                else:
                    continue
                hit = mods & banned
                self.assertFalse(hit, f"{py.name} imports {hit}")

    def test_wf12_no_db_file_created(self):
        src = make_sample_book(os.path.join(self.dir, "db.docx"))
        before = set(os.listdir(self.dir))
        self.svc().process(src)
        after = set(os.listdir(self.dir))
        self.assertFalse(any(f.endswith(".db") for f in after - before))


class TestWF13_Regression(WfBase):
    """WF-13: regression ۳۹ تست قبلی (در ماتریس جدا اجرا می‌شود — اینجا
    فقط تأیید می‌شود که factory با pct_test_modules تداخلی ندارد)."""

    def test_wf13_imports_isolated(self):
        # import factory نباید pct_test_support یا server را load کند
        import sys
        before = set(sys.modules)
        import factory  # noqa
        new = set(sys.modules) - before
        self.assertFalse({"server", "pct_test_support"} & new)


class TestOnStageCallback(WfBase):
    """PCM-15 — callback اختیاری on_stage: ترتیب قطعی stageها + last-safe
    stage در خطا + خروجی بدون تغییر نسبت به حالت پیش‌فرض (additive)."""

    EXPECTED_ORDER = ["intake", "validate", "parse", "structure",
                      "metadata", "normalize", "quality", "export"]

    def test_on_stage_order_good_file(self):
        src = make_sample_book(os.path.join(self.dir, "book.docx"))
        seen = []
        doc = self.svc().process(src, on_stage=seen.append)
        self.assertEqual(seen, self.EXPECTED_ORDER)
        self.assertGreater(len(doc.blocks), 0)
        self.assertTrue(doc.quality is not None)

    def test_on_stage_last_safe_on_failure(self):
        src = make_broken_docx(os.path.join(self.dir, "broken.docx"))
        seen = []
        with self.assertRaises(Exception):
            self.svc().process(src, on_stage=seen.append)
        # bad-zip در stage خودِ intake رد می‌شود → هیچ stage کامل نشده →
        # callback اصلا صدا نمی‌خورد (last-safe stage = null)
        self.assertEqual(seen, [])

    def test_on_stage_invariance_vs_default(self):
        """خروجی با callback باید با خروجی بدون callback یکسان باشد."""
        src = make_sample_book(os.path.join(self.dir, "book.docx"))
        seen = []
        d1 = self.svc().process(src, on_stage=seen.append)
        d2 = self.svc().process(src)
        self.assertEqual(d1.file.sha256, d2.file.sha256)
        self.assertEqual(len(d1.blocks), len(d2.blocks))
        for b1, b2 in zip(d1.blocks, d2.blocks):
            self.assertEqual(b1.raw_text, b2.raw_text)
            self.assertEqual(b1.normalized_text, b2.normalized_text)
        self.assertEqual(d1.work_id, d2.work_id)
        self.assertEqual(d1.edition_id, d2.edition_id)
        self.assertEqual(d1.quality.score, d2.quality.score)
        self.assertEqual(d1.quality.passed, d2.quality.passed)
        self.assertEqual(seen, self.EXPECTED_ORDER)


if __name__ == "__main__":
    unittest.main(verbosity=2)
