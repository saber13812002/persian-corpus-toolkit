"""backend.factory._smoke_fixtures — ساخت fixtureهای docx برای تست/دستور محلی.

(کاربری داخلی — در test_word_factory.py هم همین توابع به‌صورت fixture استفاده می‌شوند.)
"""
from __future__ import annotations

import os


def make_docx(path: str, paragraphs: list[tuple[str, str]]) -> str:
    """paragraphs: list of (style, text) — 'Heading1'.. یا 'Normal'."""
    from docx import Document
    doc = Document()
    for style, text in paragraphs:
        doc.add_paragraph(text, style=style if style != "Normal" else None)
    doc.save(path)
    return path


def make_sample_book(path: str) -> str:
    """یک کتاب نمونه کوچک با front matter، فصل‌ها، جدول و متن فارسی."""
    paras = [
        ("Normal", "عنوان: نمونه کتاب آزمون"),
        ("Normal", "نویسنده: نویسنده نمونه"),
        ("Normal", "ناشر: انتشارات نمونه"),
        ("Heading 1", "فصل اول"),
        ("Normal", "این اولین پاراگراف فصل اول است. فارسی و انگلیسی English در کنار هم."),
        ("Normal", "دومین پاراگراف با عدد ۱۴۰۵ و عدد 2026 در یک خط."),
        ("Heading 2", "بخش ۱"),
        ("Normal", "متن زیربخش با فاصلهٔ   اضافی و\tتب."),
        ("Heading 1", "فصل دوم"),
        ("Normal", "متن فصل دوم."),
    ]
    make_docx(path, paras)
    # اضافه‌کردن یک جدول به آخر
    from docx import Document
    doc = Document(path)
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "ستون1"
    t.rows[0].cells[1].text = "ستون2"
    t.rows[0].cells[2].text = "ستون3"
    t.rows[1].cells[0].text = "ارز1"
    t.rows[1].cells[1].text = "ارز2"
    t.rows[1].cells[2].text = "ارز3"
    doc.save(path)
    return path


def make_broken_docx(path: str) -> str:
    """فایلی با پسوند docx ولی محتوای نامعتبر (نه zip معتبر)."""
    with open(path, "wb") as f:
        f.write(b"PK\x03\x04" + b"\x00" * 100 + b"not a real docx")
    return path


def make_ole2_doc(path: str) -> str:
    """فایل .doc قدیمی با magic OLE2 (محتوای واقعی doc لازم نیست — فقط magic)."""
    with open(path, "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    return path


if __name__ == "__main__":
    import sys
    d = sys.argv[1] if len(sys.argv) > 1 else "."
    os.makedirs(d, exist_ok=True)
    print(make_sample_book(os.path.join(d, "sample.docx")))
    print(make_broken_docx(os.path.join(d, "broken.docx")))
    print(make_ole2_doc(os.path.join(d, "legacy.doc")))
